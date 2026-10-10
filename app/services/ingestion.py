from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}

CHUNK_SIZE = 900
CHUNK_OVERLAP = 120


def load_file(path: Path) -> list[Document]:
    """Load a supported document and attach source metadata."""
    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    if not path.is_file():
        raise ValueError(f"Expected a file: {path}")

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED:
        raise ValueError(
            f"Unsupported file type: {suffix or '[no extension]'}. "
            f"Supported types: {', '.join(sorted(SUPPORTED))}"
        )

    try:
        if suffix == ".pdf":
            documents = PyPDFLoader(str(path)).load()

        elif suffix in {".txt", ".md"}:
            documents = TextLoader(
                str(path),
                encoding="utf-8",
            ).load()

        else:  # .docx
            docx = DocxDocument(str(path))
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in docx.paragraphs
                if paragraph.text.strip()
            )

            documents = [
                Document(page_content=text, metadata={})
            ]

    except Exception as exc:
        raise RuntimeError(
            f"Failed to load '{path.name}': {exc}"
        ) from exc

    loaded = []

    for document in documents:
        if not document.page_content.strip():
            continue

        document.metadata = {
            **document.metadata,
            "source": str(path.resolve()),
            "filename": path.name,
            "file_type": suffix.lstrip("."),
        }

        loaded.append(document)

    return loaded


def chunk_documents(
    docs: Iterable[Document],
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
) -> list[Document]:
    """Split documents into overlapping chunks for RAG retrieval."""

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")

    if not 0 <= chunk_overlap < chunk_size:
        raise ValueError(
            "chunk_overlap must be non-negative and smaller than chunk_size."
        )

    documents = list(docs)

    if not documents:
        return []

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        add_start_index=True,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_documents(documents)

    source_counts: dict[str, int] = {}

    for chunk in chunks:
        source = str(chunk.metadata.get("source", "unknown"))
        index = source_counts.get(source, 0)

        chunk.metadata["chunk_index"] = index
        source_counts[source] = index + 1

    return chunks